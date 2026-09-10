from dataclasses import dataclass
from datetime import UTC, datetime
import hashlib
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from gridfs import GridFS
from openpyxl import load_workbook
from pydantic import BaseModel
from pymongo.errors import DuplicateKeyError
from pypdf import PdfReader

from app.ai.extraction import TenderExtractor
from app.models import SourceRecord


MAX_ARCHIVE_ENTRIES = 2_000
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_PDF_PAGES = 500
MAX_SPREADSHEET_ROWS = 100_000
MAX_SPREADSHEET_CELLS = 2_000_000
MAX_EXTRACTED_CHARACTERS = 2_000_000
OOXML_MAGIC = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


class UploadValidationError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ParsedDocument:
    kind: str
    text: str


class UploadExtractionResponse(BaseModel):
    upload_id: str
    file_hash: str
    records: list[SourceRecord]
    warnings: list[str]


class UploadProcessor:
    def __init__(self, database, extractor: TenderExtractor, max_bytes: int):
        self.database = database
        self.extractor = extractor
        self.max_bytes = max_bytes

    async def process(self, filename: str, content_type: str | None, content: bytes, uploader: str) -> UploadExtractionResponse:
        parsed = parse_document(filename, content, self.max_bytes)
        file_hash = hashlib.sha256(content).hexdigest()
        files = GridFS(self.database, collection="upload_files")
        uploads = self.database.uploads
        uploads.create_index("file_hash", unique=True)
        upload = uploads.find_one({"file_hash": file_hash})
        if (upload is not None and upload.get("status") == "extracted" and upload.get("records")
                and not any("Structured extraction unavailable" in warning for warning in upload.get("warnings", []))):
            return UploadExtractionResponse(
                upload_id=str(upload["_id"]),
                file_hash=file_hash,
                records=[SourceRecord.model_validate(record) for record in upload["records"]],
                warnings=upload.get("warnings", []),
            )
        if upload is None:
            file_id = files.put(
                content,
                filename=filename,
                content_type=content_type,
                metadata={"file_hash": file_hash, "uploader": uploader},
            )
            upload = {
                "file_hash": file_hash,
                "file_id": file_id,
                "filename": filename,
                "content_type": content_type,
                "kind": parsed.kind,
                "uploader": uploader,
                "created_at": datetime.now(UTC),
                "status": "stored",
            }
            try:
                upload["_id"] = uploads.insert_one(upload).inserted_id
            except DuplicateKeyError:
                files.delete(file_id)
                upload = uploads.find_one({"file_hash": file_hash})

        try:
            records, warnings = await self.extractor.extract(filename, parsed.text, file_hash)
        except Exception as error:
            uploads.update_one(
                {"_id": upload["_id"]},
                {"$set": {"status": "failed", "error": type(error).__name__, "processed_at": datetime.now(UTC)}},
            )
            raise
        uploads.update_one(
            {"_id": upload["_id"]},
            {"$set": {
                "status": "extracted",
                "record_count": len(records),
                "records": [record.model_dump(by_alias=True, mode="json") for record in records],
                "warnings": warnings,
                "processed_at": datetime.now(UTC),
            }},
        )
        return UploadExtractionResponse(
            upload_id=str(upload["_id"]),
            file_hash=file_hash,
            records=records,
            warnings=warnings,
        )


def parse_document(filename: str, content: bytes, max_bytes: int) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".docx", ".xlsx", ".pdf"}:
        raise UploadValidationError("unsupported_type", "Upload a DOCX, XLSX, or PDF file")
    if not content:
        raise UploadValidationError("empty_file", "Uploaded file is empty")
    if len(content) > max_bytes:
        raise UploadValidationError("file_too_large", f"File exceeds {max_bytes // (1024 * 1024)} MiB limit")

    if suffix == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise UploadValidationError("type_mismatch", "File extension does not match PDF content")
        text = _pdf_text(content)
    else:
        if not content.startswith(OOXML_MAGIC):
            raise UploadValidationError("type_mismatch", f"File extension does not match {suffix[1:].upper()} content")
        _validate_ooxml(content, suffix)
        text = _docx_text(content) if suffix == ".docx" else _xlsx_text(content)

    text = text.strip()
    if not text:
        message = "PDF contains no selectable text; scanned PDFs are not supported" if suffix == ".pdf" else "File contains no readable text"
        raise UploadValidationError("no_selectable_text" if suffix == ".pdf" else "no_content", message)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise UploadValidationError("content_too_large", "Extracted document content is too large")
    return ParsedDocument(suffix[1:], text)


def _validate_ooxml(content: bytes, suffix: str) -> None:
    required = "word/document.xml" if suffix == ".docx" else "xl/workbook.xml"
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise UploadValidationError("unsafe_archive", "Office archive contains too many entries")

            names = set()
            total_size = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                path = PurePosixPath(normalized)
                if path.is_absolute() or ".." in path.parts:
                    raise UploadValidationError("unsafe_archive", "Office archive contains an unsafe path")
                if entry.flag_bits & 1:
                    raise UploadValidationError("encrypted_file", "Encrypted Office files are not supported")

                names.add(normalized)
                total_size += entry.file_size
                if total_size > MAX_ARCHIVE_BYTES:
                    raise UploadValidationError("unsafe_archive", "Office archive expands beyond safe limit")
                if entry.file_size and (entry.compress_size == 0 or entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO):
                    raise UploadValidationError("unsafe_archive", "Office archive has an unsafe compression ratio")

                lower = normalized.lower()
                if lower.endswith("vbaproject.bin") or "/embeddings/" in lower or lower.endswith("oleobject.bin"):
                    raise UploadValidationError("unsafe_office_content", "Macros and embedded objects are not supported")
                if lower.endswith((".xml", ".rels")):
                    markup = archive.read(entry)
                    upper = markup.upper()
                    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                        raise UploadValidationError("unsafe_office_content", "Office file contains unsafe XML declarations")

            if required not in names:
                raise UploadValidationError("type_mismatch", f"File content is not a valid {suffix[1:].upper()} document")
    except BadZipFile as error:
        raise UploadValidationError("malformed_file", "Office file is malformed") from error


def _docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise UploadValidationError("malformed_file", "DOCX file could not be read") from error

    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append("\t".join(values))
    return "\n".join(blocks)


def _xlsx_text(content: bytes) -> str:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise UploadValidationError("malformed_file", "XLSX file could not be read") from error

    blocks = []
    row_count = 0
    cell_count = 0
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            rows = []
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                cell_count += len(row)
                if row_count > MAX_SPREADSHEET_ROWS or cell_count > MAX_SPREADSHEET_CELLS:
                    raise UploadValidationError("content_too_large", "Spreadsheet contains too many cells")
                values = [str(value).strip() if value is not None else "" for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                blocks.append(f"Worksheet: {sheet.title}\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def _pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise UploadValidationError("encrypted_file", "Encrypted PDF files are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UploadValidationError("content_too_large", "PDF contains too many pages")
        return "\n\n".join(filter(None, (page.extract_text() for page in reader.pages)))
    except UploadValidationError:
        raise
    except Exception as error:
        raise UploadValidationError("malformed_file", "PDF file could not be read") from error
