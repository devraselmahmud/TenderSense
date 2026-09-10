from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.uploads import UploadValidationError, parse_document


MAX_BYTES = 10 * 1024 * 1024


def test_parse_docx_paragraphs_and_tables():
    stream = BytesIO()
    document = Document()
    document.add_heading("Network Security Tender")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Deadline"
    table.rows[0].cells[1].text = "30 September 2026"
    document.save(stream)

    parsed = parse_document("tenders.docx", stream.getvalue(), MAX_BYTES)

    assert parsed.kind == "docx"
    assert "Network Security Tender" in parsed.text
    assert "Deadline\t30 September 2026" in parsed.text


def test_parse_visible_xlsx_rows_only():
    stream = BytesIO()
    workbook = Workbook()
    workbook.active.append(["Title", "Deadline"])
    workbook.active.append(["ERP Upgrade", "2026-09-30"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Secret Tender"])
    workbook.save(stream)

    parsed = parse_document("tenders.xlsx", stream.getvalue(), MAX_BYTES)

    assert parsed.kind == "xlsx"
    assert "Title\tDeadline" in parsed.text
    assert "ERP Upgrade\t2026-09-30" in parsed.text
    assert "Secret Tender" not in parsed.text


def test_reject_scanned_pdf():
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(stream)

    with pytest.raises(UploadValidationError, match="scanned PDFs are not supported") as error:
        parse_document("scan.pdf", stream.getvalue(), MAX_BYTES)

    assert error.value.code == "no_selectable_text"


def test_reject_extension_signature_mismatch():
    with pytest.raises(UploadValidationError) as error:
        parse_document("tenders.pdf", b"PK\x03\x04not-a-pdf", MAX_BYTES)

    assert error.value.code == "type_mismatch"


def test_reject_unsafe_ooxml_path():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../outside.xml", "<unsafe/>")

    with pytest.raises(UploadValidationError) as error:
        parse_document("tenders.docx", stream.getvalue(), MAX_BYTES)

    assert error.value.code == "unsafe_archive"


def test_reject_ooxml_entity_declaration():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY file SYSTEM "file:///etc/passwd">]><document/>')

    with pytest.raises(UploadValidationError) as error:
        parse_document("tenders.docx", stream.getvalue(), MAX_BYTES)

    assert error.value.code == "unsafe_office_content"


def test_reject_oversized_file_before_parsing():
    with pytest.raises(UploadValidationError) as error:
        parse_document("large.pdf", b"%PDF-" + b"x" * 20, 10)

    assert error.value.code == "file_too_large"
